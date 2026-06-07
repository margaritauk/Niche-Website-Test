#!/usr/bin/env python3
"""
Build the LedgerLite Shop Sales Tracker (.xlsx) — for tracking the SHOP's own sales:
which listing sells, on which platform, gross/fees/net, units, by month.
Output -> shop/Shop-Sales-Tracker.xlsx
"""
import os
from datetime import date
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

NAVY="1F3A5F"; TEAL="2C7A7B"; LIGHT="EBF2F7"; LIGHTER="F5F9FC"; WHITE="FFFFFF"
GREEN="2F855A"; RED="C53030"; GREY="718096"; BORDER="CBD5E0"
thin=Side(style="thin",color=BORDER); box=Border(left=thin,right=thin,top=thin,bottom=thin)
def hf(s=11,c=WHITE,b=True): return Font(name="Calibri",size=s,bold=b,color=c)
def bf(s=11,c="1A202C",b=False): return Font(name="Calibri",size=s,bold=b,color=c)
def fill(c): return PatternFill("solid",fgColor=c)
ctr=Alignment(horizontal="center",vertical="center",wrap_text=True)
lft=Alignment(horizontal="left",vertical="center")
rgt=Alignment(horizontal="right",vertical="center")

LISTINGS=["Short-Term Rental Tracker","Rental Property Tracker","Online Seller Tracker",
 "Freelancer Tracker","Real Estate Agent Tracker","Photographer Tracker","Cleaning Business Tracker",
 "Hair Stylist & Beauty Tracker","Rideshare & Delivery Tracker","Food Truck & Vendor Tracker",
 "Rental Investor Bundle","Self-Employed Bundle","Everything Bundle"]
PLATFORMS=["Etsy","Gumroad","Payhip","Other"]
ROWS=600; START=4

wb=Workbook()

# ---- Start Here
ws=wb.active; ws.title="Start Here"; ws.sheet_view.showGridLines=False
ws.column_dimensions["A"].width=3; ws.column_dimensions["B"].width=100
ws["B2"]="  LedgerLite — Shop Sales Tracker"; ws["B2"].font=hf(20); ws["B2"].fill=fill(NAVY)
ws["B2"].alignment=Alignment(horizontal="left",vertical="center",indent=1); ws.row_dimensions[2].height=44
tips=[("Track your own shop here — see what sells and what you actually keep.",True),
 ("1.  Each time you make a sale, add a row on the  Sales  tab: date, which listing, the platform, "
  "units, the gross price, and the fees the platform took.",False),
 ("2.  Net is calculated for you (gross − fees) — that's what you actually keep.",False),
 ("3.  The  Dashboard  shows totals, your best-selling listings, and a breakdown by platform and month.",False),
 ("Tip:  for Etsy, fees ≈ 6.5% transaction + ~3% + $0.25 payment + $0.20 listing. Gumroad ≈ 10% + fees. "
  "Enter the actual fees from your payout for accuracy.",False),
 ("This is your private business record — not tax advice. Keep your own receipts for filing.",False)]
r=4
for t,h in tips:
    c=ws[f"B{r}"]; c.value=t
    if h: c.font=hf(13,TEAL); c.fill=fill(LIGHT); c.alignment=Alignment(horizontal="left",vertical="center",indent=1); ws.row_dimensions[r].height=26
    else: c.font=bf(11); c.alignment=Alignment(horizontal="left",vertical="top",wrap_text=True); ws.row_dimensions[r].height=34
    r+=1

# ---- helper sheet for dropdowns
ref=wb.create_sheet("Lists"); ref.sheet_view.showGridLines=False
for i,v in enumerate(LISTINGS): ref.cell(row=1+i,column=1,value=v)
for i,v in enumerate(PLATFORMS): ref.cell(row=1+i,column=2,value=v)
ref.column_dimensions["A"].width=34; ref.column_dimensions["B"].width=16
LIST_R=f"Lists!$A$1:$A${len(LISTINGS)}"; PLAT_R=f"Lists!$B$1:$B${len(PLATFORMS)}"

# ---- Sales
sa=wb.create_sheet("Sales"); sa.sheet_view.showGridLines=False; sa.freeze_panes="A4"
cols=[("A",13,"Date"),("B",30,"Listing"),("C",14,"Platform"),("D",9,"Units"),
 ("E",14,"Gross $"),("F",13,"Fees $"),("G",14,"Net $"),("H",12,"Month"),("I",26,"Notes")]
for L,w,_ in cols: sa.column_dimensions[L].width=w
sa.merge_cells("A1:I1"); sa["A1"]="Sales  —  add one row per order"
sa["A1"].font=hf(14); sa["A1"].fill=fill(NAVY); sa["A1"].alignment=Alignment(horizontal="left",vertical="center",indent=1); sa.row_dimensions[1].height=30
for i,(L,w,t) in enumerate(cols):
    c=sa.cell(row=3,column=i+1,value=t); c.font=hf(11); c.fill=fill(TEAL); c.alignment=ctr; c.border=box
sa.row_dimensions[3].height=22
samples=[(date(2026,6,2),"Online Seller Tracker","Etsy",1,19.00,2.05),
 (date(2026,6,3),"Everything Bundle","Gumroad",1,69.00,7.20),
 (date(2026,6,5),"Real Estate Agent Tracker","Etsy",1,24.00,2.45)]
for i,(dte,lst,plat,u,g,f) in enumerate(samples):
    rr=START+i
    sa.cell(row=rr,column=1,value=dte); sa.cell(row=rr,column=2,value=lst); sa.cell(row=rr,column=3,value=plat)
    sa.cell(row=rr,column=4,value=u); sa.cell(row=rr,column=5,value=g); sa.cell(row=rr,column=6,value=f)
for i in range(ROWS):
    rr=START+i
    for col in range(1,10):
        c=sa.cell(row=rr,column=col); c.border=box; c.fill=fill(LIGHTER if i%2==0 else WHITE); c.font=bf(10)
        if col==1: c.number_format="yyyy-mm-dd"; c.alignment=ctr
        elif col in (5,6,7): c.number_format='$#,##0.00'; c.alignment=rgt
        elif col in (4,8): c.alignment=ctr
        else: c.alignment=lft
    sa.cell(row=rr,column=7,value=f'=IF(E{rr}="","",E{rr}-F{rr})')      # Net
    sa.cell(row=rr,column=8,value=f'=IF(A{rr}="","",TEXT(A{rr},"yyyy-mm"))')
    sa.row_dimensions[rr].height=18
dv_l=DataValidation(type="list",formula1=f"={LIST_R}",allow_blank=True)
dv_p=DataValidation(type="list",formula1=f"={PLAT_R}",allow_blank=True)
for dv in (dv_l,dv_p): dv.showErrorMessage=False; sa.add_data_validation(dv)
dv_l.add(f"B{START}:B{START+ROWS-1}"); dv_p.add(f"C{START}:C{START+ROWS-1}")

GROSS=f"Sales!$E${START}:$E${START+ROWS-1}"; FEES=f"Sales!$F${START}:$F${START+ROWS-1}"
NET=f"Sales!$G${START}:$G${START+ROWS-1}"; UNITS=f"Sales!$D${START}:$D${START+ROWS-1}"
LST=f"Sales!$B${START}:$B${START+ROWS-1}"; PLT=f"Sales!$C${START}:$C${START+ROWS-1}"
MON=f"Sales!$H${START}:$H${START+ROWS-1}"

# ---- Dashboard
db=wb.create_sheet("Dashboard"); db.sheet_view.showGridLines=False
for L,w in [("A",3),("B",30),("C",14),("D",14),("E",12),("F",4),("G",14),("H",14),("I",10)]:
    db.column_dimensions[L].width=w
db.merge_cells("B2:I2"); db["B2"]="Shop Dashboard  —  live"; db["B2"].font=hf(16); db["B2"].fill=fill(NAVY)
db["B2"].alignment=Alignment(horizontal="left",vertical="center",indent=1); db.row_dimensions[2].height=32
def kpi(col,row,label,formula,color,money=True):
    c1=db.cell(row=row,column=col,value=label); c1.font=hf(11,WHITE); c1.fill=fill(color); c1.alignment=ctr; c1.border=box
    c2=db.cell(row=row+1,column=col,value=formula); c2.font=Font(name="Calibri",size=18,bold=True,color=color)
    c2.fill=fill(LIGHTER); c2.alignment=ctr; c2.border=box
    c2.number_format='$#,##0.00' if money else '#,##0'
    db.row_dimensions[row].height=20; db.row_dimensions[row+1].height=32
kpi(2,4,"Gross sales",f'=SUM({GROSS})',NAVY)
kpi(3,4,"Fees",f'=SUM({FEES})',RED)
kpi(4,4,"Net (you keep)",f'=SUM({NET})',GREEN)
kpi(7,4,"Units sold",f'=SUM({UNITS})',TEAL,money=False)

# by listing
r=8; db.cell(row=r,column=2,value="By listing").font=hf(12,TEAL); db.cell(row=r,column=2).fill=fill(LIGHT)
for col in range(2,6): db.cell(row=r,column=col).fill=fill(LIGHT); db.cell(row=r,column=col).border=box
db.cell(row=r,column=2).alignment=Alignment(horizontal="left",indent=1); r+=1
for i,h in enumerate(["Listing","Units","Net $",""]):
    c=db.cell(row=r,column=2+i,value=h); c.font=hf(11); c.fill=fill(TEAL); c.alignment=ctr; c.border=box
r+=1
first=r
for i,lst in enumerate(LISTINGS):
    db.cell(row=r,column=2,value=lst).border=box; db.cell(row=r,column=2).font=bf(10); db.cell(row=r,column=2).alignment=lft
    db.cell(row=r,column=3,value=f'=SUMIFS({UNITS},{LST},$B{r})').number_format='#,##0'
    db.cell(row=r,column=4,value=f'=SUMIFS({NET},{LST},$B{r})').number_format='$#,##0.00'
    for col in (2,3,4): db.cell(row=r,column=col).border=box; db.cell(row=r,column=col).fill=fill(LIGHTER if i%2==0 else WHITE)
    db.cell(row=r,column=3).font=bf(10); db.cell(row=r,column=4).font=bf(10)
    r+=1

# by platform (right)
pr=8; db.cell(row=pr,column=7,value="By platform").font=hf(12,TEAL)
db.merge_cells(start_row=pr,start_column=7,end_row=pr,end_column=9)
for col in range(7,10): db.cell(row=pr,column=col).fill=fill(LIGHT); db.cell(row=pr,column=col).border=box
db.cell(row=pr,column=7).alignment=Alignment(horizontal="left",indent=1); pr+=1
for i,h in enumerate(["Platform","Net $","Units"]):
    c=db.cell(row=pr,column=7+i,value=h); c.font=hf(11); c.fill=fill(TEAL); c.alignment=ctr; c.border=box
pr+=1
for i,p in enumerate(PLATFORMS):
    db.cell(row=pr,column=7,value=p).border=box; db.cell(row=pr,column=7).font=bf(10); db.cell(row=pr,column=7).alignment=lft
    db.cell(row=pr,column=8,value=f'=SUMIFS({NET},{PLT},$G{pr})').number_format='$#,##0.00'
    db.cell(row=pr,column=9,value=f'=SUMIFS({UNITS},{PLT},$G{pr})').number_format='#,##0'
    for col in (7,8,9): db.cell(row=pr,column=col).border=box; db.cell(row=pr,column=col).fill=fill(LIGHTER if i%2==0 else WHITE)
    pr+=1
# by month (right, below platform)
pr+=1; db.cell(row=pr,column=7,value="By month (2026)").font=hf(12,TEAL)
db.merge_cells(start_row=pr,start_column=7,end_row=pr,end_column=9)
for col in range(7,10): db.cell(row=pr,column=col).fill=fill(LIGHT); db.cell(row=pr,column=col).border=box
db.cell(row=pr,column=7).alignment=Alignment(horizontal="left",indent=1); pr+=1
for i,h in enumerate(["Month","Net $","Units"]):
    c=db.cell(row=pr,column=7+i,value=h); c.font=hf(11); c.fill=fill(TEAL); c.alignment=ctr; c.border=box
pr+=1
for i in range(1,13):
    m=f"2026-{i:02d}"
    db.cell(row=pr,column=7,value=m).border=box; db.cell(row=pr,column=7).font=bf(10); db.cell(row=pr,column=7).alignment=ctr
    db.cell(row=pr,column=8,value=f'=SUMIFS({NET},{MON},$G{pr})').number_format='$#,##0.00'
    db.cell(row=pr,column=9,value=f'=SUMIFS({UNITS},{MON},$G{pr})').number_format='#,##0'
    for col in (7,8,9): db.cell(row=pr,column=col).border=box; db.cell(row=pr,column=col).fill=fill(LIGHTER if i%2==0 else WHITE)
    pr+=1

wb.active=0
out=os.path.join(os.path.dirname(os.path.abspath(__file__)),"shop","Shop-Sales-Tracker.xlsx")
wb.save(out); print("wrote",os.path.relpath(out))
