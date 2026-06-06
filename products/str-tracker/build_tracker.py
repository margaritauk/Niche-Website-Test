#!/usr/bin/env python3
"""
Build the Short-Term Rental (Airbnb/VRBO) Bookkeeping & Tax Tracker .xlsx

Generates a polished, formula-driven Excel workbook that:
  - Logs income & expenses per property
  - Auto-maps expenses to IRS Schedule E lines
  - Produces a year-end, tax-ready Schedule E summary per property
  - Shows a live dashboard (net profit, by-property, by-month)

This same engine builds sibling variants (long-term landlord, single-trade
bookkeeping) by swapping the CATEGORIES / labels below.

Usage:  python3 build_tracker.py
Output: Short-Term-Rental-Tracker.xlsx
"""

from datetime import date
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, NamedStyle
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

# ---------------------------------------------------------------- palette
NAVY      = "1F3A5F"   # headers
TEAL      = "2C7A7B"   # accents
LIGHT     = "EBF2F7"   # zebra / panels
LIGHTER   = "F5F9FC"
GREEN     = "2F855A"
RED       = "C53030"
GREY      = "718096"
WHITE     = "FFFFFF"
BORDER_C  = "CBD5E0"

thin = Side(style="thin", color=BORDER_C)
box  = Border(left=thin, right=thin, top=thin, bottom=thin)

def hfont(sz=11, color=WHITE, bold=True): return Font(name="Calibri", size=sz, bold=bold, color=color)
def bfont(sz=11, color="1A202C", bold=False): return Font(name="Calibri", size=sz, bold=bold, color=color)
def fill(c): return PatternFill("solid", fgColor=c)

center = Alignment(horizontal="center", vertical="center", wrap_text=True)
left   = Alignment(horizontal="left",   vertical="center", wrap_text=True)
right  = Alignment(horizontal="right",  vertical="center")
topl   = Alignment(horizontal="left",   vertical="top",    wrap_text=True)

# ---------------------------------------------------------------- data model
MAX_PROPS = 5
TX_ROWS   = 1000          # transaction capacity
TX_START  = 4            # first data row on Transactions (1=title,2=spacer,3=header)

INCOME_CATEGORIES = [
    "Rental income (nightly)",
    "Cleaning fee collected",
    "Other guest fees",
]

# (label, Schedule E line)  -- Schedule E Part I, lines 5-19
EXPENSE_CATEGORIES = [
    ("Advertising",                  "5  Advertising"),
    ("Auto and travel",              "6  Auto and travel"),
    ("Cleaning and maintenance",     "7  Cleaning and maintenance"),
    ("Commissions / platform fees",  "8  Commissions"),
    ("Insurance",                    "9  Insurance"),
    ("Legal and professional fees",  "10 Legal & professional"),
    ("Management fees",              "11 Management fees"),
    ("Mortgage interest (banks)",    "12 Mortgage interest"),
    ("Other interest",               "13 Other interest"),
    ("Repairs",                      "14 Repairs"),
    ("Supplies / consumables",       "15 Supplies"),
    ("Taxes (incl. occupancy tax)",  "16 Taxes"),
    ("Utilities",                    "17 Utilities"),
    ("Depreciation",                 "18 Depreciation"),
    ("Other (software, fees, misc.)","19 Other"),
]

ALL_CATEGORIES = INCOME_CATEGORIES + [c[0] for c in EXPENSE_CATEGORIES]

wb = Workbook()

# currency style
cur = NamedStyle(name="cur")
cur.number_format = '$#,##0.00'
cur.alignment = right
wb.add_named_style(cur)

# ================================================================ 1. START HERE
ws = wb.active
ws.title = "Start Here"
ws.sheet_view.showGridLines = False
ws.column_dimensions["A"].width = 3
ws.column_dimensions["B"].width = 100
ws.column_dimensions["C"].width = 3

def banner(ws, row, text, sub=None):
    ws.merge_cells(f"B{row}:B{row}")
    c = ws[f"B{row}"]; c.value = text
    c.font = hfont(20); c.fill = fill(NAVY); c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[row].height = 46
    if sub:
        ws[f"B{row+1}"] = sub
        ws[f"B{row+1}"].font = bfont(11, GREY, bold=False)
        ws[f"B{row+1}"].alignment = left
        ws.row_dimensions[row+1].height = 22

banner(ws, 2, "  Short-Term Rental  —  Bookkeeping & Tax Tracker",
       "  Airbnb / VRBO / direct-booking hosts  ·  IRS Schedule E ready  ·  works in Excel & Google Sheets")

steps = [
    ("How this works", None, True),
    ("1.  Open the  Setup  tab and type your property names (up to 5). Everything else "
     "links to these names automatically.", None, False),
    ("2.  Each time money moves, add one row on the  Transactions  tab. Pick the property, "
     "choose Income or Expense, and pick a category from the dropdown. That's it.", None, False),
    ("3.  The  Schedule E Summary  tab fills itself in — every expense is mapped to the correct "
     "IRS Schedule E line, per property. Hand it straight to your accountant or copy the totals "
     "into your tax software.", None, False),
    ("4.  The  Dashboard  tab updates live: net profit, income vs. expenses, totals by property "
     "and by month.", None, False),
    ("Tips", None, True),
    ("•  Categories live on the  Categories  tab. The dropdowns read from there, so you never "
     "mistype a category.", None, False),
    ("•  Log the GROSS booking as income and the platform's cut as a 'Commissions / platform fees' "
     "expense — that's how it should appear on Schedule E.", None, False),
    ("•  Occupancy / lodging taxes you collect and remit go under 'Taxes (incl. occupancy tax)'.", None, False),
    ("•  Duplicate the file per tax year (e.g. 'STR Tracker 2026').", None, False),
    ("Disclaimer", None, True),
    ("This template is a bookkeeping organizer, not tax, legal, or accounting advice. Schedule E "
     "line mapping is provided for convenience — confirm your specific situation with a qualified "
     "tax professional.", None, False),
]
r = 5
for text, _, header in steps:
    cell = ws[f"B{r}"]; cell.value = text
    if header:
        cell.font = hfont(13, TEAL); cell.fill = fill(LIGHT)
        cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws.row_dimensions[r].height = 26
    else:
        cell.font = bfont(11); cell.alignment = topl
        ws.row_dimensions[r].height = 34
    r += 1

# ================================================================ 2. SETUP
sp = wb.create_sheet("Setup")
sp.sheet_view.showGridLines = False
sp.column_dimensions["A"].width = 3
sp.column_dimensions["B"].width = 34
sp.column_dimensions["C"].width = 26
sp.column_dimensions["D"].width = 26
sp.column_dimensions["E"].width = 22

sp.merge_cells("B2:E2")
sp["B2"] = "Property Setup"
sp["B2"].font = hfont(16); sp["B2"].fill = fill(NAVY)
sp["B2"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
sp.row_dimensions[2].height = 34
sp.merge_cells("B3:E3")
sp["B3"] = "Type each property's name below. These names feed every dropdown and report."
sp["B3"].font = bfont(10, GREY); sp["B3"].alignment = left

hdr = ["Property name", "Address (optional)", "Platform", "Purchase price (optional)"]
for i, h in enumerate(hdr):
    c = sp.cell(row=5, column=2+i, value=h)
    c.font = hfont(11); c.fill = fill(TEAL); c.alignment = center; c.border = box
sp.row_dimensions[5].height = 24

placeholders = ["Beach Cottage", "Downtown Loft", "Mountain Cabin", "", ""]
for i in range(MAX_PROPS):
    row = 6 + i
    nm = sp.cell(row=row, column=2, value=placeholders[i] or None)
    nm.font = bfont(11, bold=True)
    for col in range(2, 6):
        cc = sp.cell(row=row, column=col)
        cc.border = box
        cc.fill = fill(LIGHTER if i % 2 == 0 else WHITE)
        if col == 5:
            cc.style = "cur"
    sp.cell(row=row, column=2).alignment = left
    sp.row_dimensions[row].height = 22

PROP_RANGE = f"Setup!$B$6:$B${6+MAX_PROPS-1}"

# ================================================================ 3. CATEGORIES
cat = wb.create_sheet("Categories")
cat.sheet_view.showGridLines = False
cat.column_dimensions["A"].width = 3
cat.column_dimensions["B"].width = 36
cat.column_dimensions["C"].width = 30
cat.column_dimensions["D"].width = 4
cat.column_dimensions["E"].width = 30

cat.merge_cells("B2:C2")
cat["B2"] = "Categories  &  Schedule E mapping"
cat["B2"].font = hfont(14); cat["B2"].fill = fill(NAVY)
cat["B2"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
cat.row_dimensions[2].height = 30

# Expense table (col B = category, col C = schedule E line)
cat["B4"] = "Expense category"; cat["C4"] = "Schedule E line"
for cc in ("B4", "C4"):
    cat[cc].font = hfont(11); cat[cc].fill = fill(TEAL); cat[cc].alignment = center; cat[cc].border = box
for i, (label, line) in enumerate(EXPENSE_CATEGORIES):
    row = 5 + i
    a = cat.cell(row=row, column=2, value=label); a.border = box; a.font = bfont(10); a.alignment = left
    b = cat.cell(row=row, column=3, value=line);  b.border = box; b.font = bfont(10); b.alignment = left
    f = LIGHTER if i % 2 == 0 else WHITE
    a.fill = fill(f); b.fill = fill(f)

# Income list (col E)
cat["E4"] = "Income category"
cat["E4"].font = hfont(11); cat["E4"].fill = fill(GREEN); cat["E4"].alignment = center; cat["E4"].border = box
for i, label in enumerate(INCOME_CATEGORIES):
    c = cat.cell(row=5+i, column=5, value=label); c.border = box; c.font = bfont(10); c.alignment = left
    c.fill = fill(LIGHTER if i % 2 == 0 else WHITE)

# Master list used by the Transactions dropdown (all categories), placed in col H (hidden-ish)
cat["H1"] = "All categories (dropdown source)"; cat["H1"].font = bfont(9, GREY)
for i, label in enumerate(ALL_CATEGORIES):
    cat.cell(row=2+i, column=8, value=label).font = bfont(9, GREY)
cat.column_dimensions["F"].width = 3
cat.column_dimensions["G"].width = 3
cat.column_dimensions["H"].width = 34
ALLCAT_RANGE = f"Categories!$H$2:$H${1+len(ALL_CATEGORIES)}"

# also keep an income-only and expense-only range for type validation help
cat["J1"] = "Type"; cat.cell(row=2, column=10, value="Income"); cat.cell(row=3, column=10, value="Expense")
cat.column_dimensions["I"].width = 3
cat.column_dimensions["J"].width = 12
TYPE_RANGE = "Categories!$J$2:$J$3"

# ================================================================ 4. TRANSACTIONS
tx = wb.create_sheet("Transactions")
tx.sheet_view.showGridLines = False
tx.freeze_panes = "A4"
cols = [("A", 13, "Date"), ("B", 22, "Property"), ("C", 13, "Type"),
        ("D", 30, "Category"), ("E", 34, "Description"), ("F", 15, "Amount"),
        ("G", 16, "Platform / payee"), ("H", 14, "Month"), ("I", 26, "Notes")]
for letter, width, _ in cols:
    tx.column_dimensions[letter].width = width

tx.merge_cells("A1:I1")
tx["A1"] = "Transactions  —  add one row per income or expense"
tx["A1"].font = hfont(14); tx["A1"].fill = fill(NAVY)
tx["A1"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
tx.row_dimensions[1].height = 30

for i, (letter, width, title) in enumerate(cols):
    c = tx.cell(row=3, column=i+1, value=title)
    c.font = hfont(11); c.fill = fill(TEAL); c.alignment = center; c.border = box
tx.row_dimensions[3].height = 24

# sample rows so the buyer sees how it works
samples = [
    (date(2026,5,2),  "Beach Cottage", "Income",  "Rental income (nightly)",   "Booking #A12 — 4 nights",        1240.00, "Airbnb"),
    (date(2026,5,2),  "Beach Cottage", "Income",  "Cleaning fee collected",    "Booking #A12 cleaning fee",        95.00, "Airbnb"),
    (date(2026,5,2),  "Beach Cottage", "Expense", "Commissions / platform fees","Airbnb host service fee",          46.50, "Airbnb"),
    (date(2026,5,5),  "Beach Cottage", "Expense", "Cleaning and maintenance",  "Turnover cleaning",                85.00, "Sparkle Co"),
    (date(2026,5,9),  "Downtown Loft", "Income",  "Rental income (nightly)",   "Booking #V7 — 3 nights",          870.00, "VRBO"),
    (date(2026,5,11), "Downtown Loft", "Expense", "Supplies / consumables",    "Coffee, paper goods, soap",        38.20, "Costco"),
    (date(2026,5,15), "Beach Cottage", "Expense", "Utilities",                 "Electric + internet",             142.00, "Utility Co"),
    (date(2026,5,20), "Downtown Loft", "Expense", "Taxes (incl. occupancy tax)","Occupancy tax remitted (May)",    78.30, "State DOR"),
]
for i, (d, prop, typ, c, desc, amt, plat) in enumerate(samples):
    row = TX_START + i
    tx.cell(row=row, column=1, value=d)
    tx.cell(row=row, column=2, value=prop)
    tx.cell(row=row, column=3, value=typ)
    tx.cell(row=row, column=4, value=c)
    tx.cell(row=row, column=5, value=desc)
    tx.cell(row=row, column=6, value=amt)
    tx.cell(row=row, column=7, value=plat)

# style + Month formula for all capacity rows
for i in range(TX_ROWS):
    row = TX_START + i
    for col in range(1, 10):
        cc = tx.cell(row=row, column=col)
        cc.border = box
        cc.fill = fill(LIGHTER if i % 2 == 0 else WHITE)
        cc.font = bfont(10)
        if col in (1,):
            cc.number_format = "yyyy-mm-dd"; cc.alignment = center
        elif col == 6:
            cc.number_format = '$#,##0.00'; cc.alignment = right
        elif col == 8:
            cc.alignment = center
        else:
            cc.alignment = left
    # Month helper: YYYY-MM as text, only when a date exists
    tx.cell(row=row, column=8,
            value=f'=IF(A{row}="","",TEXT(A{row},"yyyy-mm"))')
    tx.row_dimensions[row].height = 18

TX_PROP = f"Transactions!$B${TX_START}:$B${TX_START+TX_ROWS-1}"
TX_TYPE = f"Transactions!$C${TX_START}:$C${TX_START+TX_ROWS-1}"
TX_CAT  = f"Transactions!$D${TX_START}:$D${TX_START+TX_ROWS-1}"
TX_AMT  = f"Transactions!$F${TX_START}:$F${TX_START+TX_ROWS-1}"
TX_MON  = f"Transactions!$H${TX_START}:$H${TX_START+TX_ROWS-1}"

# data validations
dv_prop = DataValidation(type="list", formula1=f"={PROP_RANGE}", allow_blank=True)
dv_type = DataValidation(type="list", formula1=f"={TYPE_RANGE}", allow_blank=True)
dv_cat  = DataValidation(type="list", formula1=f"={ALLCAT_RANGE}", allow_blank=True)
for dv in (dv_prop, dv_type, dv_cat):
    dv.showErrorMessage = False
    tx.add_data_validation(dv)
dv_prop.add(f"B{TX_START}:B{TX_START+TX_ROWS-1}")
dv_type.add(f"C{TX_START}:C{TX_START+TX_ROWS-1}")
dv_cat.add(f"D{TX_START}:D{TX_START+TX_ROWS-1}")

# ================================================================ 5. SCHEDULE E SUMMARY
se = wb.create_sheet("Schedule E Summary")
se.sheet_view.showGridLines = False
se.column_dimensions["A"].width = 3
se.column_dimensions["B"].width = 34   # line label
for i in range(MAX_PROPS):
    se.column_dimensions[get_column_letter(3+i)].width = 16
se.column_dimensions[get_column_letter(3+MAX_PROPS)].width = 16  # total col

se.merge_cells("B2:" + get_column_letter(3+MAX_PROPS) + "2")
se["B2"] = "Schedule E Summary  (Part I)  —  auto-calculated per property"
se["B2"].font = hfont(14); se["B2"].fill = fill(NAVY)
se["B2"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
se.row_dimensions[2].height = 30

# header row of property names (linked to Setup)
se["B4"] = "Schedule E line"
se["B4"].font = hfont(11); se["B4"].fill = fill(TEAL); se["B4"].alignment = left; se["B4"].border = box
for i in range(MAX_PROPS):
    col = 3 + i
    c = se.cell(row=4, column=col, value=f"=IF({_prop_ref(i)}=\"\",\"(property {i+1})\",{_prop_ref(i)})" if False else None)
    # link directly to Setup cell
    c.value = f"=Setup!$B${6+i}"
    c.font = hfont(11); c.fill = fill(TEAL); c.alignment = center; c.border = box
totcol = 3 + MAX_PROPS
c = se.cell(row=4, column=totcol, value="TOTAL")
c.font = hfont(11); c.fill = fill(NAVY); c.alignment = center; c.border = box
se.row_dimensions[4].height = 22

def sumifs_for(catlabel, prop_col_setup_row):
    # sum expense amounts where property matches the setup cell and category matches
    return (f'=SUMIFS({TX_AMT},{TX_PROP},Setup!$B${prop_col_setup_row},'
            f'{TX_CAT},"{catlabel}",{TX_TYPE},"Expense")')

# --- INCOME block ---
row = 5
se.cell(row=row, column=2, value="INCOME").font = hfont(11, NAVY)
se.cell(row=row, column=2).fill = fill(LIGHT); se.cell(row=row, column=2).border = box
for i in range(MAX_PROPS+1):
    cc = se.cell(row=row, column=3+i); cc.fill = fill(LIGHT); cc.border = box
row += 1
income_first = row
for label in INCOME_CATEGORIES:
    se.cell(row=row, column=2, value=label).font = bfont(10)
    se.cell(row=row, column=2).alignment = left; se.cell(row=row, column=2).border = box
    for i in range(MAX_PROPS):
        col = 3+i
        f = (f'=SUMIFS({TX_AMT},{TX_PROP},Setup!$B${6+i},'
             f'{TX_CAT},"{label}",{TX_TYPE},"Income")')
        cc = se.cell(row=row, column=col, value=f)
        cc.number_format = '$#,##0.00'; cc.border = box
    # row total
    cc = se.cell(row=row, column=totcol,
                 value=f"=SUM({get_column_letter(3)}{row}:{get_column_letter(2+MAX_PROPS)}{row})")
    cc.number_format = '$#,##0.00'; cc.border = box; cc.font = bfont(10, bold=True)
    row += 1
income_last = row - 1
# total income
se.cell(row=row, column=2, value="Total income (Line 3)").font = bfont(10, GREEN, bold=True)
se.cell(row=row, column=2).border = box; se.cell(row=row, column=2).fill = fill(LIGHTER)
for i in range(MAX_PROPS+1):
    col = 3+i
    L = get_column_letter(col)
    cc = se.cell(row=row, column=col, value=f"=SUM({L}{income_first}:{L}{income_last})")
    cc.number_format = '$#,##0.00'; cc.border = box; cc.font = bfont(10, GREEN, bold=True); cc.fill = fill(LIGHTER)
income_total_row = row
row += 2

# --- EXPENSE block ---
se.cell(row=row, column=2, value="EXPENSES").font = hfont(11, NAVY)
se.cell(row=row, column=2).fill = fill(LIGHT); se.cell(row=row, column=2).border = box
for i in range(MAX_PROPS+1):
    cc = se.cell(row=row, column=3+i); cc.fill = fill(LIGHT); cc.border = box
row += 1
exp_first = row
for label, line in EXPENSE_CATEGORIES:
    se.cell(row=row, column=2, value=line).font = bfont(10)
    se.cell(row=row, column=2).alignment = left; se.cell(row=row, column=2).border = box
    for i in range(MAX_PROPS):
        col = 3+i
        cc = se.cell(row=row, column=col, value=sumifs_for(label, 6+i))
        cc.number_format = '$#,##0.00'; cc.border = box
    L0, L1 = get_column_letter(3), get_column_letter(2+MAX_PROPS)
    cc = se.cell(row=row, column=totcol, value=f"=SUM({L0}{row}:{L1}{row})")
    cc.number_format = '$#,##0.00'; cc.border = box; cc.font = bfont(10, bold=True)
    row += 1
exp_last = row - 1
# total expenses (line 20)
se.cell(row=row, column=2, value="Total expenses (Line 20)").font = bfont(10, RED, bold=True)
se.cell(row=row, column=2).border = box; se.cell(row=row, column=2).fill = fill(LIGHTER)
for i in range(MAX_PROPS+1):
    col = 3+i; L = get_column_letter(col)
    cc = se.cell(row=row, column=col, value=f"=SUM({L}{exp_first}:{L}{exp_last})")
    cc.number_format = '$#,##0.00'; cc.border = box; cc.font = bfont(10, RED, bold=True); cc.fill = fill(LIGHTER)
exp_total_row = row
row += 1
# net income (line 21-ish)
se.cell(row=row, column=2, value="NET INCOME / (LOSS)").font = hfont(11, WHITE)
se.cell(row=row, column=2).fill = fill(TEAL); se.cell(row=row, column=2).border = box
for i in range(MAX_PROPS+1):
    col = 3+i; L = get_column_letter(col)
    cc = se.cell(row=row, column=col, value=f"={L}{income_total_row}-{L}{exp_total_row}")
    cc.number_format = '$#,##0.00;[Red]($#,##0.00)'; cc.border = box
    cc.font = hfont(11, WHITE); cc.fill = fill(TEAL)
se.row_dimensions[row].height = 22

# ================================================================ 6. DASHBOARD
db = wb.create_sheet("Dashboard")
db.sheet_view.showGridLines = False
db.column_dimensions["A"].width = 3
db.column_dimensions["B"].width = 30
db.column_dimensions["C"].width = 18
db.column_dimensions["D"].width = 6
db.column_dimensions["E"].width = 24
db.column_dimensions["F"].width = 16
db.column_dimensions["G"].width = 16
db.column_dimensions["H"].width = 16

db.merge_cells("B2:H2")
db["B2"] = "Dashboard  —  live overview"
db["B2"].font = hfont(16); db["B2"].fill = fill(NAVY)
db["B2"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
db.row_dimensions[2].height = 34

def kpi(anchor_col, row, label, formula, color):
    c1 = db.cell(row=row, column=anchor_col, value=label)
    c1.font = hfont(11, WHITE); c1.fill = fill(color); c1.alignment = center; c1.border = box
    c2 = db.cell(row=row+1, column=anchor_col, value=formula)
    c2.font = Font(name="Calibri", size=18, bold=True, color=color)
    c2.fill = fill(LIGHTER); c2.alignment = center; c2.border = box
    c2.number_format = '$#,##0.00;[Red]($#,##0.00)'
    db.row_dimensions[row].height = 20
    db.row_dimensions[row+1].height = 34

tot_income = f'=SUMIFS({TX_AMT},{TX_TYPE},"Income")'
tot_exp    = f'=SUMIFS({TX_AMT},{TX_TYPE},"Expense")'
kpi(2, 4, "Total income",   tot_income, GREEN)
kpi(3, 4, "Total expenses", tot_exp,    RED)
db.merge_cells("E4:F4"); db.merge_cells("E5:F5")
db["E4"] = "Net profit / (loss)"
db["E4"].font = hfont(11, WHITE); db["E4"].fill = fill(TEAL); db["E4"].alignment = center; db["E4"].border = box
db["E5"] = f'=({tot_income[1:]})-({tot_exp[1:]})'
db["E5"].font = Font(name="Calibri", size=20, bold=True, color=TEAL)
db["E5"].fill = fill(LIGHT); db["E5"].alignment = center; db["E5"].border = box
db["E5"].number_format = '$#,##0.00;[Red]($#,##0.00)'
db["F4"].border = box; db["F5"].border = box

# --- by property ---
r = 8
db.cell(row=r, column=2, value="By property").font = hfont(12, TEAL)
db.cell(row=r, column=2).fill = fill(LIGHT)
db.cell(row=r, column=2).alignment = Alignment(horizontal="left", indent=1)
for col in range(2, 6):
    db.cell(row=r, column=col).fill = fill(LIGHT); db.cell(row=r, column=col).border = box
r += 1
for i, h in enumerate(["Property", "Income", "Expenses", "Net"]):
    c = db.cell(row=r, column=2+i, value=h)
    c.font = hfont(11); c.fill = fill(TEAL); c.alignment = center; c.border = box
r += 1
for i in range(MAX_PROPS):
    setup_cell = f"Setup!$B${6+i}"
    db.cell(row=r, column=2, value=f"={setup_cell}").border = box
    db.cell(row=r, column=2).font = bfont(10, bold=True); db.cell(row=r, column=2).alignment = left
    inc = f'=SUMIFS({TX_AMT},{TX_PROP},{setup_cell},{TX_TYPE},"Income")'
    exp = f'=SUMIFS({TX_AMT},{TX_PROP},{setup_cell},{TX_TYPE},"Expense")'
    db.cell(row=r, column=3, value=inc).number_format = '$#,##0.00'
    db.cell(row=r, column=4, value=exp).number_format = '$#,##0.00'
    db.cell(row=r, column=5, value=f"=C{r}-D{r}").number_format = '$#,##0.00;[Red]($#,##0.00)'
    for col in range(2, 6):
        cc = db.cell(row=r, column=col); cc.border = box
        cc.fill = fill(LIGHTER if i % 2 == 0 else WHITE)
        if col >= 3: cc.font = bfont(10)
    r += 1
# totals
db.cell(row=r, column=2, value="All properties").font = bfont(10, NAVY, bold=True)
db.cell(row=r, column=2).fill = fill(LIGHT); db.cell(row=r, column=2).border = box; db.cell(row=r, column=2).alignment = left
db.cell(row=r, column=3, value=f"=SUM(C{r-MAX_PROPS}:C{r-1})").number_format = '$#,##0.00'
db.cell(row=r, column=4, value=f"=SUM(D{r-MAX_PROPS}:D{r-1})").number_format = '$#,##0.00'
db.cell(row=r, column=5, value=f"=SUM(E{r-MAX_PROPS}:E{r-1})").number_format = '$#,##0.00;[Red]($#,##0.00)'
for col in range(2, 6):
    cc = db.cell(row=r, column=col); cc.border = box; cc.fill = fill(LIGHT); cc.font = bfont(10, NAVY, bold=True)

# --- by month (right side) ---
months = [f"2026-{m:02d}" for m in range(1, 13)]
mr = 8
db.cell(row=mr, column=6, value="By month (2026)").font = hfont(12, TEAL)
db.cell(row=mr, column=6).fill = fill(LIGHT)
db.merge_cells(start_row=mr, start_column=6, end_row=mr, end_column=8)
for col in range(6, 9):
    db.cell(row=mr, column=col).fill = fill(LIGHT); db.cell(row=mr, column=col).border = box
db.cell(row=mr, column=6).alignment = Alignment(horizontal="left", indent=1)
mr += 1
for i, h in enumerate(["Month", "Income", "Expenses"]):
    c = db.cell(row=mr, column=6+i, value=h)
    c.font = hfont(11); c.fill = fill(TEAL); c.alignment = center; c.border = box
mr += 1
for i, m in enumerate(months):
    db.cell(row=mr, column=6, value=m).alignment = center; db.cell(row=mr, column=6).font = bfont(10)
    inc = f'=SUMIFS({TX_AMT},{TX_MON},"{m}",{TX_TYPE},"Income")'
    exp = f'=SUMIFS({TX_AMT},{TX_MON},"{m}",{TX_TYPE},"Expense")'
    db.cell(row=mr, column=7, value=inc).number_format = '$#,##0.00'
    db.cell(row=mr, column=8, value=exp).number_format = '$#,##0.00'
    for col in range(6, 9):
        cc = db.cell(row=mr, column=col); cc.border = box
        cc.fill = fill(LIGHTER if i % 2 == 0 else WHITE)
        if col >= 7: cc.font = bfont(10)
    mr += 1

# ---------------------------------------------------------------- save
wb.active = 0
out = "Short-Term-Rental-Tracker.xlsx"
wb.save(out)
print("wrote", out)
