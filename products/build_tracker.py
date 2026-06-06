#!/usr/bin/env python3
"""
Build formula-driven Bookkeeping & Tax Tracker workbooks (.xlsx).

One engine, multiple products. Each VARIANT below produces a polished workbook that:
  - Logs income & expenses per property
  - Auto-maps expenses to IRS Schedule E lines
  - Produces a year-end, tax-ready Schedule E summary per property
  - Shows a live dashboard (net profit, by-property, by-month)

Usage:
  python3 build_tracker.py            # build all variants
  python3 build_tracker.py str        # build one variant by key

Add a new product = add a dict to VARIANTS. No other code changes needed.
"""

import sys
import os
from datetime import date
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, NamedStyle
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

# ---------------------------------------------------------------- palette
NAVY    = "1F3A5F"; TEAL = "2C7A7B"; LIGHT = "EBF2F7"; LIGHTER = "F5F9FC"
GREEN   = "2F855A"; RED  = "C53030"; GREY  = "718096"; WHITE   = "FFFFFF"
BORDER_C = "CBD5E0"

thin = Side(style="thin", color=BORDER_C)
box  = Border(left=thin, right=thin, top=thin, bottom=thin)

def hfont(sz=11, color=WHITE, bold=True): return Font(name="Calibri", size=sz, bold=bold, color=color)
def bfont(sz=11, color="1A202C", bold=False): return Font(name="Calibri", size=sz, bold=bold, color=color)
def fill(c): return PatternFill("solid", fgColor=c)

center = Alignment(horizontal="center", vertical="center", wrap_text=True)
left   = Alignment(horizontal="left",   vertical="center", wrap_text=True)
right  = Alignment(horizontal="right",  vertical="center")
topl   = Alignment(horizontal="left",   vertical="top",    wrap_text=True)

MAX_PROPS = 5
TX_ROWS   = 1000
TX_START  = 4

# Standard IRS Schedule E Part I expense lines (5-19). Labels vary per variant.
def sched_e(labels):
    """labels: list of 15 strings, in Schedule E line order 5..19."""
    lines = ["5  Advertising","6  Auto and travel","7  Cleaning and maintenance",
             "8  Commissions","9  Insurance","10 Legal & professional","11 Management fees",
             "12 Mortgage interest","13 Other interest","14 Repairs","15 Supplies",
             "16 Taxes","17 Utilities","18 Depreciation","19 Other"]
    return list(zip(labels, lines))

# ================================================================ VARIANTS
VARIANTS = {
    "str": {
        "dir": "str-tracker",
        "output": "Short-Term-Rental-Tracker.xlsx",
        "title": "  Short-Term Rental  —  Bookkeeping & Tax Tracker",
        "subtitle": "  Airbnb / VRBO / direct-booking hosts  ·  IRS Schedule E ready  ·  works in Excel & Google Sheets",
        "income": ["Rental income (nightly)", "Cleaning fee collected", "Other guest fees"],
        "expenses": sched_e([
            "Advertising","Auto and travel","Cleaning and maintenance","Commissions / platform fees",
            "Insurance","Legal and professional fees","Management fees","Mortgage interest (banks)",
            "Other interest","Repairs","Supplies / consumables","Taxes (incl. occupancy tax)",
            "Utilities","Depreciation","Other (software, fees, misc.)"]),
        "placeholders": ["Beach Cottage", "Downtown Loft", "Mountain Cabin", "", ""],
        "setup_col3": "Platform",
        "payee_header": "Platform / payee",
        "tip_gross": ("Log the GROSS booking as income and the platform's cut as a "
                      "'Commissions / platform fees' expense — that's how it should appear on Schedule E."),
        "tip_extra": ("Occupancy / lodging taxes you collect and remit go under "
                      "'Taxes (incl. occupancy tax)'."),
        "noun": "short-term rental",
        "audience_line": "If you host on Airbnb, VRBO, or take direct bookings,",
        "log_right": [
            "<b>Log the GROSS payout as income, and the platform's cut as an expense.</b> Example: a guest "
            "pays $1,240. Enter $1,240 as <i>Rental income (nightly)</i>, then enter Airbnb's $46.50 host "
            "fee as a separate <i>Commissions / platform fees</i> expense. That's exactly how it should "
            "appear on Schedule E.",
            "<b>Cleaning fees you collect</b> are income (<i>Cleaning fee collected</i>); what you pay your "
            "cleaner is an expense (<i>Cleaning and maintenance</i>).",
            "<b>Occupancy / lodging taxes</b> you collect and remit go under <i>Taxes (incl. occupancy tax)</i>.",
            "<b>One transaction per row.</b> Don't combine. The more granular you are, the cleaner your "
            "tax summary.",
        ],
        "samples": [
            (date(2026,5,2),  "Beach Cottage", "Income",  "Rental income (nightly)",    "Booking #A12 — 4 nights",     1240.00, "Airbnb"),
            (date(2026,5,2),  "Beach Cottage", "Income",  "Cleaning fee collected",     "Booking #A12 cleaning fee",     95.00, "Airbnb"),
            (date(2026,5,2),  "Beach Cottage", "Expense", "Commissions / platform fees","Airbnb host service fee",       46.50, "Airbnb"),
            (date(2026,5,5),  "Beach Cottage", "Expense", "Cleaning and maintenance",   "Turnover cleaning",             85.00, "Sparkle Co"),
            (date(2026,5,9),  "Downtown Loft", "Income",  "Rental income (nightly)",    "Booking #V7 — 3 nights",       870.00, "VRBO"),
            (date(2026,5,11), "Downtown Loft", "Expense", "Supplies / consumables",     "Coffee, paper goods, soap",     38.20, "Costco"),
            (date(2026,5,15), "Beach Cottage", "Expense", "Utilities",                  "Electric + internet",          142.00, "Utility Co"),
            (date(2026,5,20), "Downtown Loft", "Expense", "Taxes (incl. occupancy tax)","Occupancy tax remitted (May)",   78.30, "State DOR"),
        ],
    },
    "landlord": {
        "dir": "landlord-tracker",
        "output": "Rental-Property-Tracker.xlsx",
        "title": "  Rental Property  —  Bookkeeping & Tax Tracker",
        "subtitle": "  Long-term landlords  ·  IRS Schedule E ready  ·  works in Excel & Google Sheets",
        "income": ["Rent received", "Late fees", "Other income (pet/parking/laundry)"],
        "expenses": sched_e([
            "Advertising","Auto and travel","Cleaning and maintenance","Commissions / leasing fees",
            "Insurance","Legal and professional fees","Management fees","Mortgage interest (banks)",
            "Other interest","Repairs","Supplies","Property taxes","Utilities","Depreciation",
            "Other (HOA, misc.)"]),
        "placeholders": ["123 Oak St", "456 Maple Ave", "Unit 7B", "", ""],
        "setup_col3": "Lease type",
        "payee_header": "Tenant / payee",
        "tip_gross": ("Log each rent payment as 'Rent received'. A security deposit is NOT income — "
                      "only log it as income if and when you actually keep it."),
        "tip_extra": ("Your annual property tax goes under 'Property taxes'; HOA dues go under "
                      "'Other (HOA, misc.)'."),
        "noun": "rental property",
        "audience_line": "If you rent out long-term residential property,",
        "log_right": [
            "<b>Log each rent payment as income</b> (<i>Rent received</i>). A late fee you charge is "
            "income too (<i>Late fees</i>).",
            "<b>A security deposit is NOT income.</b> Don't log it when you receive it — only log it as "
            "income if and when you actually keep part of it.",
            "<b>On your mortgage, only the INTEREST portion is a Schedule E expense</b> "
            "(<i>Mortgage interest (banks)</i>) — not the principal. Your lender's year-end 1098 shows the "
            "interest figure.",
            "<b>Property tax</b> goes under <i>Property taxes</i>; <b>HOA dues</b> under "
            "<i>Other (HOA, misc.)</i>. <b>One transaction per row</b> keeps the summary clean.",
        ],
        "samples": [
            (date(2026,5,1),  "123 Oak St",     "Income",  "Rent received",            "May rent",                  1650.00, "Tenant: J. Lee"),
            (date(2026,5,1),  "456 Maple Ave",  "Income",  "Rent received",            "May rent",                  1875.00, "Tenant: M. Diaz"),
            (date(2026,5,3),  "123 Oak St",     "Expense", "Repairs",                  "Kitchen faucet replacement",  145.00, "HandyPro"),
            (date(2026,5,5),  "123 Oak St",     "Expense", "Mortgage interest (banks)","May interest portion",        612.00, "BigBank"),
            (date(2026,5,10), "456 Maple Ave",  "Expense", "Management fees",          "8% management fee",           150.00, "PM Co"),
            (date(2026,5,15), "123 Oak St",     "Expense", "Insurance",                "Landlord policy (monthly)",    78.00, "Insurer"),
            (date(2026,5,20), "456 Maple Ave",  "Expense", "Property taxes",           "Q2 property tax",             410.00, "County"),
            (date(2026,5,28), "123 Oak St",     "Income",  "Late fees",                "Late rent fee",                50.00, "Tenant: J. Lee"),
        ],
    },
    "seller": {
        "dir": "seller-tracker",
        "output": "Online-Seller-Bookkeeping-Tracker.xlsx",
        "form": "Schedule C",
        "entity": "Shop",
        "entity_lower": "shop",
        "entity_plural_lower": "shops",
        "summary_sheet": "Schedule C Summary",
        "summary_part": "",
        "income_line_label": "Total income (Line 1)",
        "expense_total_label": "Total expenses (incl. COGS)",
        "net_label": "NET PROFIT / (LOSS)",
        "setup_col2": "Notes (optional)",
        "setup_col3": "Platform",
        "setup_col4": "Start date (optional)",
        "setup_col4_currency": False,
        "title": "  Online Seller  —  Bookkeeping & Tax Tracker",
        "subtitle": "  Etsy / Amazon / Shopify / craft-fair sellers  ·  IRS Schedule C ready  ·  works in Excel & Google Sheets",
        "income": ["Product sales", "Shipping income collected", "Other income"],
        "expenses": [
            ("Advertising & promotion",          "8  Advertising"),
            ("Car & truck",                      "9  Car & truck"),
            ("Commissions & marketplace fees",   "10 Commissions & fees"),
            ("Contract labor",                   "11 Contract labor"),
            ("Cost of goods sold (materials)",   "4  Cost of goods sold"),
            ("Insurance (business)",             "15 Insurance"),
            ("Legal & professional",             "17 Legal & professional"),
            ("Office expense",                   "18 Office expense"),
            ("Rent or lease",                    "20 Rent or lease"),
            ("Repairs & maintenance",            "21 Repairs & maintenance"),
            ("Shipping & postage",               "27 Other (shipping)"),
            ("Software & subscriptions",         "27 Other (software)"),
            ("Supplies & packaging",             "22 Supplies"),
            ("Taxes & licenses",                 "23 Taxes & licenses"),
            ("Travel & meals",                   "24 Travel & meals"),
            ("Utilities (phone/internet)",       "25 Utilities"),
        ],
        "placeholders": ["Etsy Shop", "Amazon Store", "Craft Fairs", "", ""],
        "payee_header": "Customer / payee",
        "tip_gross": ("Log gross sales as income and the marketplace's cut (Etsy/Amazon/PayPal/Square "
                      "fees) as a 'Commissions & marketplace fees' expense — don't net them together."),
        "tip_extra": ("Materials and inventory you buy to make or resell products go under "
                      "'Cost of goods sold (materials)'; packaging goes under 'Supplies & packaging'."),
        "noun": "online shop",
        "audience_line": "If you sell on Etsy, Amazon, Shopify, or at markets,",
        "log_right": [
            "<b>Log gross sales as income</b> (<i>Product sales</i>), and the marketplace's cut as an "
            "expense (<i>Commissions & marketplace fees</i>). Etsy/Amazon/payment fees are real, "
            "deductible costs — track them separately so your numbers are right.",
            "<b>Materials and inventory</b> you buy to make or resell go under <i>Cost of goods sold "
            "(materials)</i>. Boxes, mailers and labels go under <i>Supplies & packaging</i> or "
            "<i>Shipping & postage</i>.",
            "<b>Shipping you charge the customer</b> is income (<i>Shipping income collected</i>); "
            "<b>postage you pay</b> is an expense (<i>Shipping & postage</i>).",
            "<b>One transaction per row.</b> Run more than one shop? Give each its own name on the "
            "Setup tab — the summary breaks results out per shop while still totaling your business.",
        ],
        "samples": [
            (date(2026,5,2),  "Etsy Shop",    "Income",  "Product sales",                  "Orders #1001–1010",           980.00, "Etsy"),
            (date(2026,5,2),  "Etsy Shop",    "Expense", "Commissions & marketplace fees", "Etsy transaction + payment",   88.20, "Etsy"),
            (date(2026,5,4),  "Etsy Shop",    "Expense", "Cost of goods sold (materials)", "Yarn + beads restock",        145.00, "Supplier"),
            (date(2026,5,6),  "Etsy Shop",    "Expense", "Shipping & postage",             "May postage labels",           62.40, "USPS"),
            (date(2026,5,9),  "Amazon Store", "Income",  "Product sales",                  "FBA payout",                 1320.00, "Amazon"),
            (date(2026,5,9),  "Amazon Store", "Expense", "Commissions & marketplace fees", "Amazon referral + FBA fees",  264.00, "Amazon"),
            (date(2026,5,15), "Etsy Shop",    "Expense", "Software & subscriptions",       "Canva + email tool",           29.00, "SaaS"),
            (date(2026,5,20), "Craft Fairs",  "Income",  "Product sales",                  "Spring market booth",         410.00, "Square"),
        ],
    },
    "freelancer": {
        "dir": "freelancer-tracker",
        "output": "Freelancer-Self-Employed-Bookkeeping-Tracker.xlsx",
        "form": "Schedule C",
        "entity": "Client",
        "entity_lower": "client",
        "entity_plural_lower": "clients",
        "summary_sheet": "Schedule C Summary",
        "summary_part": "",
        "income_line_label": "Total income (Line 1)",
        "expense_total_label": "Total expenses (Line 28)",
        "net_label": "NET PROFIT / (LOSS)",
        "setup_col2": "Notes (optional)",
        "setup_col3": "Type (client / overhead)",
        "setup_col4": "Start date (optional)",
        "setup_col4_currency": False,
        "title": "  Freelancer & Self-Employed  —  Bookkeeping & Tax Tracker",
        "subtitle": "  Freelancers · 1099 contractors · solo service businesses  ·  IRS Schedule C ready  ·  works in Excel & Google Sheets",
        "income": ["Client payment / invoice", "Other income"],
        "expenses": [
            ("Advertising & marketing",          "8  Advertising"),
            ("Car & truck",                      "9  Car & truck"),
            ("Platform & payment fees",          "10 Commissions & fees"),
            ("Contract labor / subcontractors",  "11 Contract labor"),
            ("Insurance (business)",             "15 Insurance"),
            ("Legal & professional",             "17 Legal & professional"),
            ("Office expense",                   "18 Office expense"),
            ("Rent or lease",                    "20 Rent or lease"),
            ("Repairs & maintenance",            "21 Repairs & maintenance"),
            ("Supplies",                         "22 Supplies"),
            ("Taxes & licenses",                 "23 Taxes & licenses"),
            ("Travel & meals",                   "24 Travel & meals"),
            ("Software & subscriptions",         "27 Other (software)"),
            ("Education & professional dev",     "27 Other (education)"),
            ("Utilities (phone/internet)",       "25 Utilities"),
        ],
        "placeholders": ["(General / overhead)", "Client: Acme Co", "Client: Beta LLC", "", ""],
        "payee_header": "Paid to / from",
        "tip_gross": ("Log each client payment as 'Client payment / invoice' and tag it to that client "
                      "on the Setup tab — the dashboard then shows you revenue per client."),
        "tip_extra": ("Tag client-specific costs to that client; put general overhead (software, phone, "
                      "insurance) under '(General / overhead)' so it isn't charged against one client."),
        "noun": "freelance / self-employed",
        "audience_line": "If you freelance, contract, or run a solo service business,",
        "log_right": [
            "<b>Log every client payment as income</b> (<i>Client payment / invoice</i>) and tag it to "
            "that client on the Setup tab — the dashboard then shows revenue per client, your most useful "
            "number.",
            "<b>Tag client-specific costs to that client</b> (e.g. a subcontractor hired for one project). "
            "Put general overhead — software, phone, insurance — under <b>(General / overhead)</b> so it "
            "isn't double-counted against a single client.",
            "<b>Set aside for quarterly taxes.</b> This tracker shows your <i>profit</i>; as a 1099 / "
            "self-employed worker you generally owe estimated tax on it every quarter. A common rule of "
            "thumb is to park ~25–30% of profit aside — ask your tax pro for your exact number.",
            "<b>Mileage and the home-office deduction are figured separately</b> by the IRS (standard "
            "mileage rate / Form 8829). Keep a mileage log and your home-office details — this sheet "
            "covers your direct business expenses.",
        ],
        "samples": [
            (date(2026,5,3),  "Client: Acme Co",      "Income",  "Client payment / invoice",        "Invoice #210 — design work",   2200.00, "Acme Co"),
            (date(2026,5,6),  "(General / overhead)", "Expense", "Software & subscriptions",        "Adobe + Figma",                  74.00, "Adobe/Figma"),
            (date(2026,5,10), "Client: Beta LLC",     "Income",  "Client payment / invoice",        "Invoice #211 — retainer",      1500.00, "Beta LLC"),
            (date(2026,5,12), "(General / overhead)", "Expense", "Platform & payment fees",         "Stripe processing fees",         55.50, "Stripe"),
            (date(2026,5,14), "Client: Acme Co",      "Expense", "Contract labor / subcontractors", "Freelance copywriter",          400.00, "Subcontractor"),
            (date(2026,5,18), "(General / overhead)", "Expense", "Utilities (phone/internet)",      "Phone + internet (biz %)",       60.00, "Telecom"),
            (date(2026,5,22), "(General / overhead)", "Expense", "Travel & meals",                  "Client lunch + parking",         48.00, "Cafe"),
            (date(2026,5,28), "Client: Beta LLC",     "Income",  "Client payment / invoice",        "Invoice #214 — extra hours",    650.00, "Beta LLC"),
        ],
    },
}

# ================================================================ BUILDER
def build(cfg):
    INCOME_CATEGORIES = cfg["income"]
    EXPENSE_CATEGORIES = cfg["expenses"]
    ALL_CATEGORIES = INCOME_CATEGORIES + [c[0] for c in EXPENSE_CATEGORIES]

    # parameterizable labels (defaults keep the Schedule E / "property" products unchanged)
    FORM      = cfg.get("form", "Schedule E")
    ENTITY    = cfg.get("entity", "Property")            # singular, capitalized
    ENT_LOW   = cfg.get("entity_lower", "property")
    ENT_PLUR  = cfg.get("entity_plural_lower", "properties")
    SUM_SHEET = cfg.get("summary_sheet", f"{FORM} Summary")
    SUM_PART  = cfg.get("summary_part", "(Part I)")
    INC_LABEL = cfg.get("income_line_label", "Total income (Line 3)")
    EXP_LABEL = cfg.get("expense_total_label", "Total expenses (Line 20)")
    NET_LABEL = cfg.get("net_label", "NET INCOME / (LOSS)")
    SET_COL2  = cfg.get("setup_col2", "Address (optional)")
    SET_COL4  = cfg.get("setup_col4", "Purchase price (optional)")

    wb = Workbook()

    cur = NamedStyle(name="cur")
    cur.number_format = '$#,##0.00'; cur.alignment = right
    wb.add_named_style(cur)

    # -------------------------------------------------- 1. START HERE
    ws = wb.active; ws.title = "Start Here"; ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 3; ws.column_dimensions["B"].width = 100
    ws.column_dimensions["C"].width = 3

    c = ws["B2"]; c.value = cfg["title"]
    c.font = hfont(20); c.fill = fill(NAVY)
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[2].height = 46
    ws["B3"] = cfg["subtitle"]; ws["B3"].font = bfont(11, GREY); ws["B3"].alignment = left
    ws.row_dimensions[3].height = 22

    steps = [
        ("How this works", True),
        (f"1.  Open the  Setup  tab and type your {ENT_LOW} names (up to 5). Everything else "
         "links to these names automatically.", False),
        (f"2.  Each time money moves, add one row on the  Transactions  tab. Pick the {ENT_LOW}, "
         "choose Income or Expense, and pick a category from the dropdown. That's it.", False),
        (f"3.  The  {SUM_SHEET}  tab fills itself in — every expense is mapped to the correct "
         f"IRS {FORM} line, per {ENT_LOW}. Hand it straight to your accountant or copy the totals "
         "into your tax software.", False),
        (f"4.  The  Dashboard  tab updates live: net profit, income vs. expenses, totals by {ENT_LOW} "
         "and by month.", False),
        ("Tips", True),
        ("•  Categories live on the  Categories  tab. The dropdowns read from there, so you never "
         "mistype a category.", False),
        ("•  " + cfg["tip_gross"], False),
        ("•  " + cfg["tip_extra"], False),
        ("•  Duplicate the file per tax year (e.g. '" + cfg["title"].strip().split("  ")[0] + " 2026').", False),
        ("Disclaimer", True),
        (f"This template is a bookkeeping organizer, not tax, legal, or accounting advice. {FORM} "
         "line mapping is provided for convenience — confirm your specific situation with a qualified "
         "tax professional.", False),
    ]
    r = 5
    for text, header in steps:
        cell = ws[f"B{r}"]; cell.value = text
        if header:
            cell.font = hfont(13, TEAL); cell.fill = fill(LIGHT)
            cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
            ws.row_dimensions[r].height = 26
        else:
            cell.font = bfont(11); cell.alignment = topl
            ws.row_dimensions[r].height = 34
        r += 1

    # -------------------------------------------------- 2. SETUP
    sp = wb.create_sheet("Setup"); sp.sheet_view.showGridLines = False
    sp.column_dimensions["A"].width = 3; sp.column_dimensions["B"].width = 34
    sp.column_dimensions["C"].width = 26; sp.column_dimensions["D"].width = 26
    sp.column_dimensions["E"].width = 22
    sp.merge_cells("B2:E2"); sp["B2"] = f"{ENTITY} Setup"
    sp["B2"].font = hfont(16); sp["B2"].fill = fill(NAVY)
    sp["B2"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    sp.row_dimensions[2].height = 34
    sp.merge_cells("B3:E3")
    sp["B3"] = f"Type each {ENT_LOW}'s name below. These names feed every dropdown and report."
    sp["B3"].font = bfont(10, GREY); sp["B3"].alignment = left
    for i, h in enumerate([f"{ENTITY} name", SET_COL2, cfg["setup_col3"], SET_COL4]):
        cc = sp.cell(row=5, column=2+i, value=h)
        cc.font = hfont(11); cc.fill = fill(TEAL); cc.alignment = center; cc.border = box
    sp.row_dimensions[5].height = 24
    col4_cur = cfg.get("setup_col4_currency", True)
    for i in range(MAX_PROPS):
        row = 6 + i
        nm = sp.cell(row=row, column=2, value=cfg["placeholders"][i] or None)
        nm.font = bfont(11, bold=True)
        for col in range(2, 6):
            cc = sp.cell(row=row, column=col); cc.border = box
            cc.fill = fill(LIGHTER if i % 2 == 0 else WHITE)
            if col == 5 and col4_cur: cc.style = "cur"
        sp.cell(row=row, column=2).alignment = left
        sp.row_dimensions[row].height = 22

    # -------------------------------------------------- 3. CATEGORIES
    cat = wb.create_sheet("Categories"); cat.sheet_view.showGridLines = False
    cat.column_dimensions["A"].width = 3; cat.column_dimensions["B"].width = 36
    cat.column_dimensions["C"].width = 30; cat.column_dimensions["D"].width = 4
    cat.column_dimensions["E"].width = 30
    cat.merge_cells("B2:C2"); cat["B2"] = f"Categories  &  {FORM} mapping"
    cat["B2"].font = hfont(14); cat["B2"].fill = fill(NAVY)
    cat["B2"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    cat.row_dimensions[2].height = 30
    cat["B4"] = "Expense category"; cat["C4"] = f"{FORM} line"
    for cc in ("B4", "C4"):
        cat[cc].font = hfont(11); cat[cc].fill = fill(TEAL); cat[cc].alignment = center; cat[cc].border = box
    for i, (label, line) in enumerate(EXPENSE_CATEGORIES):
        row = 5 + i; f = LIGHTER if i % 2 == 0 else WHITE
        a = cat.cell(row=row, column=2, value=label); a.border = box; a.font = bfont(10); a.alignment = left; a.fill = fill(f)
        b = cat.cell(row=row, column=3, value=line);  b.border = box; b.font = bfont(10); b.alignment = left; b.fill = fill(f)
    cat["E4"] = "Income category"
    cat["E4"].font = hfont(11); cat["E4"].fill = fill(GREEN); cat["E4"].alignment = center; cat["E4"].border = box
    for i, label in enumerate(INCOME_CATEGORIES):
        c2 = cat.cell(row=5+i, column=5, value=label); c2.border = box; c2.font = bfont(10); c2.alignment = left
        c2.fill = fill(LIGHTER if i % 2 == 0 else WHITE)
    cat["H1"] = "All categories (dropdown source)"; cat["H1"].font = bfont(9, GREY)
    for i, label in enumerate(ALL_CATEGORIES):
        cat.cell(row=2+i, column=8, value=label).font = bfont(9, GREY)
    cat.column_dimensions["F"].width = 3; cat.column_dimensions["G"].width = 3
    cat.column_dimensions["H"].width = 34
    ALLCAT_RANGE = f"Categories!$H$2:$H${1+len(ALL_CATEGORIES)}"
    cat["J1"] = "Type"; cat.cell(row=2, column=10, value="Income"); cat.cell(row=3, column=10, value="Expense")
    cat.column_dimensions["I"].width = 3; cat.column_dimensions["J"].width = 12
    TYPE_RANGE = "Categories!$J$2:$J$3"
    PROP_RANGE = f"Setup!$B$6:$B${6+MAX_PROPS-1}"

    # -------------------------------------------------- 4. TRANSACTIONS
    tx = wb.create_sheet("Transactions"); tx.sheet_view.showGridLines = False
    tx.freeze_panes = "A4"
    cols = [("A", 13, "Date"), ("B", 22, ENTITY), ("C", 13, "Type"),
            ("D", 30, "Category"), ("E", 34, "Description"), ("F", 15, "Amount"),
            ("G", 16, cfg["payee_header"]), ("H", 14, "Month"), ("I", 26, "Notes")]
    for letter, width, _ in cols:
        tx.column_dimensions[letter].width = width
    tx.merge_cells("A1:I1")
    tx["A1"] = "Transactions  —  add one row per income or expense"
    tx["A1"].font = hfont(14); tx["A1"].fill = fill(NAVY)
    tx["A1"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    tx.row_dimensions[1].height = 30
    for i, (letter, width, title) in enumerate(cols):
        cc = tx.cell(row=3, column=i+1, value=title)
        cc.font = hfont(11); cc.fill = fill(TEAL); cc.alignment = center; cc.border = box
    tx.row_dimensions[3].height = 24
    for i, (d, prop, typ, c2, desc, amt, plat) in enumerate(cfg["samples"]):
        row = TX_START + i
        tx.cell(row=row, column=1, value=d)
        tx.cell(row=row, column=2, value=prop)
        tx.cell(row=row, column=3, value=typ)
        tx.cell(row=row, column=4, value=c2)
        tx.cell(row=row, column=5, value=desc)
        tx.cell(row=row, column=6, value=amt)
        tx.cell(row=row, column=7, value=plat)
    for i in range(TX_ROWS):
        row = TX_START + i
        for col in range(1, 10):
            cc = tx.cell(row=row, column=col); cc.border = box
            cc.fill = fill(LIGHTER if i % 2 == 0 else WHITE); cc.font = bfont(10)
            if col == 1:
                cc.number_format = "yyyy-mm-dd"; cc.alignment = center
            elif col == 6:
                cc.number_format = '$#,##0.00'; cc.alignment = right
            elif col == 8:
                cc.alignment = center
            else:
                cc.alignment = left
        tx.cell(row=row, column=8, value=f'=IF(A{row}="","",TEXT(A{row},"yyyy-mm"))')
        tx.row_dimensions[row].height = 18

    TX_PROP = f"Transactions!$B${TX_START}:$B${TX_START+TX_ROWS-1}"
    TX_TYPE = f"Transactions!$C${TX_START}:$C${TX_START+TX_ROWS-1}"
    TX_CAT  = f"Transactions!$D${TX_START}:$D${TX_START+TX_ROWS-1}"
    TX_AMT  = f"Transactions!$F${TX_START}:$F${TX_START+TX_ROWS-1}"
    TX_MON  = f"Transactions!$H${TX_START}:$H${TX_START+TX_ROWS-1}"

    dv_prop = DataValidation(type="list", formula1=f"={PROP_RANGE}", allow_blank=True)
    dv_type = DataValidation(type="list", formula1=f"={TYPE_RANGE}", allow_blank=True)
    dv_cat  = DataValidation(type="list", formula1=f"={ALLCAT_RANGE}", allow_blank=True)
    for dv in (dv_prop, dv_type, dv_cat):
        dv.showErrorMessage = False; tx.add_data_validation(dv)
    dv_prop.add(f"B{TX_START}:B{TX_START+TX_ROWS-1}")
    dv_type.add(f"C{TX_START}:C{TX_START+TX_ROWS-1}")
    dv_cat.add(f"D{TX_START}:D{TX_START+TX_ROWS-1}")

    # -------------------------------------------------- 5. TAX SUMMARY
    se = wb.create_sheet(SUM_SHEET); se.sheet_view.showGridLines = False
    se.column_dimensions["A"].width = 3; se.column_dimensions["B"].width = 34
    for i in range(MAX_PROPS+1):
        se.column_dimensions[get_column_letter(3+i)].width = 16
    totcol = 3 + MAX_PROPS
    se.merge_cells("B2:" + get_column_letter(totcol) + "2")
    _part = f"  {SUM_PART}" if SUM_PART else ""
    se["B2"] = f"{SUM_SHEET}{_part}  —  auto-calculated per {ENT_LOW}"
    se["B2"].font = hfont(14); se["B2"].fill = fill(NAVY)
    se["B2"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    se.row_dimensions[2].height = 30
    se["B4"] = f"{FORM} line"
    se["B4"].font = hfont(11); se["B4"].fill = fill(TEAL); se["B4"].alignment = left; se["B4"].border = box
    for i in range(MAX_PROPS):
        cc = se.cell(row=4, column=3+i, value=f"=Setup!$B${6+i}")
        cc.font = hfont(11); cc.fill = fill(TEAL); cc.alignment = center; cc.border = box
    cc = se.cell(row=4, column=totcol, value="TOTAL")
    cc.font = hfont(11); cc.fill = fill(NAVY); cc.alignment = center; cc.border = box
    se.row_dimensions[4].height = 22

    # INCOME block
    row = 5
    se.cell(row=row, column=2, value="INCOME").font = hfont(11, NAVY)
    se.cell(row=row, column=2).fill = fill(LIGHT); se.cell(row=row, column=2).border = box
    for i in range(MAX_PROPS+1):
        cc = se.cell(row=row, column=3+i); cc.fill = fill(LIGHT); cc.border = box
    row += 1; income_first = row
    for label in INCOME_CATEGORIES:
        se.cell(row=row, column=2, value=label).font = bfont(10)
        se.cell(row=row, column=2).alignment = left; se.cell(row=row, column=2).border = box
        for i in range(MAX_PROPS):
            f = (f'=SUMIFS({TX_AMT},{TX_PROP},Setup!$B${6+i},{TX_CAT},"{label}",{TX_TYPE},"Income")')
            cc = se.cell(row=row, column=3+i, value=f); cc.number_format = '$#,##0.00'; cc.border = box
        cc = se.cell(row=row, column=totcol,
                     value=f"=SUM({get_column_letter(3)}{row}:{get_column_letter(2+MAX_PROPS)}{row})")
        cc.number_format = '$#,##0.00'; cc.border = box; cc.font = bfont(10, bold=True)
        row += 1
    income_last = row - 1
    se.cell(row=row, column=2, value=INC_LABEL).font = bfont(10, GREEN, bold=True)
    se.cell(row=row, column=2).border = box; se.cell(row=row, column=2).fill = fill(LIGHTER)
    for i in range(MAX_PROPS+1):
        L = get_column_letter(3+i)
        cc = se.cell(row=row, column=3+i, value=f"=SUM({L}{income_first}:{L}{income_last})")
        cc.number_format = '$#,##0.00'; cc.border = box; cc.font = bfont(10, GREEN, bold=True); cc.fill = fill(LIGHTER)
    income_total_row = row; row += 2

    # EXPENSE block
    se.cell(row=row, column=2, value="EXPENSES").font = hfont(11, NAVY)
    se.cell(row=row, column=2).fill = fill(LIGHT); se.cell(row=row, column=2).border = box
    for i in range(MAX_PROPS+1):
        cc = se.cell(row=row, column=3+i); cc.fill = fill(LIGHT); cc.border = box
    row += 1; exp_first = row
    for label, line in EXPENSE_CATEGORIES:
        se.cell(row=row, column=2, value=line).font = bfont(10)
        se.cell(row=row, column=2).alignment = left; se.cell(row=row, column=2).border = box
        for i in range(MAX_PROPS):
            f = (f'=SUMIFS({TX_AMT},{TX_PROP},Setup!$B${6+i},{TX_CAT},"{label}",{TX_TYPE},"Expense")')
            cc = se.cell(row=row, column=3+i, value=f); cc.number_format = '$#,##0.00'; cc.border = box
        L0, L1 = get_column_letter(3), get_column_letter(2+MAX_PROPS)
        cc = se.cell(row=row, column=totcol, value=f"=SUM({L0}{row}:{L1}{row})")
        cc.number_format = '$#,##0.00'; cc.border = box; cc.font = bfont(10, bold=True)
        row += 1
    exp_last = row - 1
    se.cell(row=row, column=2, value=EXP_LABEL).font = bfont(10, RED, bold=True)
    se.cell(row=row, column=2).border = box; se.cell(row=row, column=2).fill = fill(LIGHTER)
    for i in range(MAX_PROPS+1):
        L = get_column_letter(3+i)
        cc = se.cell(row=row, column=3+i, value=f"=SUM({L}{exp_first}:{L}{exp_last})")
        cc.number_format = '$#,##0.00'; cc.border = box; cc.font = bfont(10, RED, bold=True); cc.fill = fill(LIGHTER)
    exp_total_row = row; row += 1
    se.cell(row=row, column=2, value=NET_LABEL).font = hfont(11, WHITE)
    se.cell(row=row, column=2).fill = fill(TEAL); se.cell(row=row, column=2).border = box
    for i in range(MAX_PROPS+1):
        L = get_column_letter(3+i)
        cc = se.cell(row=row, column=3+i, value=f"={L}{income_total_row}-{L}{exp_total_row}")
        cc.number_format = '$#,##0.00;[Red]($#,##0.00)'; cc.border = box
        cc.font = hfont(11, WHITE); cc.fill = fill(TEAL)
    se.row_dimensions[row].height = 22

    # -------------------------------------------------- 6. DASHBOARD
    db = wb.create_sheet("Dashboard"); db.sheet_view.showGridLines = False
    db.column_dimensions["A"].width = 3; db.column_dimensions["B"].width = 30
    db.column_dimensions["C"].width = 18; db.column_dimensions["D"].width = 6
    db.column_dimensions["E"].width = 24; db.column_dimensions["F"].width = 16
    db.column_dimensions["G"].width = 16; db.column_dimensions["H"].width = 16
    db.merge_cells("B2:H2"); db["B2"] = "Dashboard  —  live overview"
    db["B2"].font = hfont(16); db["B2"].fill = fill(NAVY)
    db["B2"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    db.row_dimensions[2].height = 34

    def kpi(anchor_col, rr, label, formula, color):
        c1 = db.cell(row=rr, column=anchor_col, value=label)
        c1.font = hfont(11, WHITE); c1.fill = fill(color); c1.alignment = center; c1.border = box
        c2 = db.cell(row=rr+1, column=anchor_col, value=formula)
        c2.font = Font(name="Calibri", size=18, bold=True, color=color)
        c2.fill = fill(LIGHTER); c2.alignment = center; c2.border = box
        c2.number_format = '$#,##0.00;[Red]($#,##0.00)'
        db.row_dimensions[rr].height = 20; db.row_dimensions[rr+1].height = 34

    tot_income = f'=SUMIFS({TX_AMT},{TX_TYPE},"Income")'
    tot_exp    = f'=SUMIFS({TX_AMT},{TX_TYPE},"Expense")'
    kpi(2, 4, "Total income",   tot_income, GREEN)
    kpi(3, 4, "Total expenses", tot_exp,    RED)
    db.merge_cells("E4:F4"); db.merge_cells("E5:F5")
    db["E4"] = "Net profit / (loss)"
    db["E4"].font = hfont(11, WHITE); db["E4"].fill = fill(TEAL); db["E4"].alignment = center; db["E4"].border = box
    db["E5"] = f'=({tot_income[1:]})-({tot_exp[1:]})'
    db["E5"].font = Font(name="Calibri", size=20, bold=True, color=TEAL)
    db["E5"].fill = fill(LIGHT); db["E5"].alignment = center; db["E5"].border = box
    db["E5"].number_format = '$#,##0.00;[Red]($#,##0.00)'
    db["F4"].border = box; db["F5"].border = box

    r = 8
    db.cell(row=r, column=2, value=f"By {ENT_LOW}").font = hfont(12, TEAL)
    db.cell(row=r, column=2).fill = fill(LIGHT)
    db.cell(row=r, column=2).alignment = Alignment(horizontal="left", indent=1)
    for col in range(2, 6):
        db.cell(row=r, column=col).fill = fill(LIGHT); db.cell(row=r, column=col).border = box
    r += 1
    for i, h in enumerate([ENTITY, "Income", "Expenses", "Net"]):
        cc = db.cell(row=r, column=2+i, value=h)
        cc.font = hfont(11); cc.fill = fill(TEAL); cc.alignment = center; cc.border = box
    r += 1
    for i in range(MAX_PROPS):
        setup_cell = f"Setup!$B${6+i}"
        db.cell(row=r, column=2, value=f"={setup_cell}").border = box
        db.cell(row=r, column=2).font = bfont(10, bold=True); db.cell(row=r, column=2).alignment = left
        db.cell(row=r, column=3, value=f'=SUMIFS({TX_AMT},{TX_PROP},{setup_cell},{TX_TYPE},"Income")').number_format = '$#,##0.00'
        db.cell(row=r, column=4, value=f'=SUMIFS({TX_AMT},{TX_PROP},{setup_cell},{TX_TYPE},"Expense")').number_format = '$#,##0.00'
        db.cell(row=r, column=5, value=f"=C{r}-D{r}").number_format = '$#,##0.00;[Red]($#,##0.00)'
        for col in range(2, 6):
            cc = db.cell(row=r, column=col); cc.border = box
            cc.fill = fill(LIGHTER if i % 2 == 0 else WHITE)
            if col >= 3: cc.font = bfont(10)
        r += 1
    db.cell(row=r, column=2, value=f"All {ENT_PLUR}").font = bfont(10, NAVY, bold=True)
    db.cell(row=r, column=2).fill = fill(LIGHT); db.cell(row=r, column=2).border = box; db.cell(row=r, column=2).alignment = left
    db.cell(row=r, column=3, value=f"=SUM(C{r-MAX_PROPS}:C{r-1})").number_format = '$#,##0.00'
    db.cell(row=r, column=4, value=f"=SUM(D{r-MAX_PROPS}:D{r-1})").number_format = '$#,##0.00'
    db.cell(row=r, column=5, value=f"=SUM(E{r-MAX_PROPS}:E{r-1})").number_format = '$#,##0.00;[Red]($#,##0.00)'
    for col in range(2, 6):
        cc = db.cell(row=r, column=col); cc.border = box; cc.fill = fill(LIGHT); cc.font = bfont(10, NAVY, bold=True)

    months = [f"2026-{m:02d}" for m in range(1, 13)]
    mr = 8
    db.cell(row=mr, column=6, value="By month (2026)").font = hfont(12, TEAL)
    db.cell(row=mr, column=6).fill = fill(LIGHT)
    db.merge_cells(start_row=mr, start_column=6, end_row=mr, end_column=8)
    for col in range(6, 9):
        db.cell(row=mr, column=col).fill = fill(LIGHT); db.cell(row=mr, column=col).border = box
    db.cell(row=mr, column=6).alignment = Alignment(horizontal="left", indent=1)
    mr += 1
    for i, h in enumerate(["Month", "Income", "Expenses"]):
        cc = db.cell(row=mr, column=6+i, value=h)
        cc.font = hfont(11); cc.fill = fill(TEAL); cc.alignment = center; cc.border = box
    mr += 1
    for i, m in enumerate(months):
        db.cell(row=mr, column=6, value=m).alignment = center; db.cell(row=mr, column=6).font = bfont(10)
        db.cell(row=mr, column=7, value=f'=SUMIFS({TX_AMT},{TX_MON},"{m}",{TX_TYPE},"Income")').number_format = '$#,##0.00'
        db.cell(row=mr, column=8, value=f'=SUMIFS({TX_AMT},{TX_MON},"{m}",{TX_TYPE},"Expense")').number_format = '$#,##0.00'
        for col in range(6, 9):
            cc = db.cell(row=mr, column=col); cc.border = box
            cc.fill = fill(LIGHTER if i % 2 == 0 else WHITE)
            if col >= 7: cc.font = bfont(10)
        mr += 1

    wb.active = 0
    out = os.path.join(os.path.dirname(os.path.abspath(__file__)), cfg["dir"], cfg["output"])
    os.makedirs(os.path.dirname(out), exist_ok=True)
    wb.save(out)
    print("wrote", out)


if __name__ == "__main__":
    keys = sys.argv[1:] or list(VARIANTS)
    for k in keys:
        if k not in VARIANTS:
            print("unknown variant:", k, "(choose from:", ", ".join(VARIANTS) + ")"); continue
        build(VARIANTS[k])
